from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

from eidon_os.application.personal_intelligence import PersonalSnapshot
from eidon_os.domain.activity import Activity


class ExcelDashboardExporter:
    def export(
        self,
        snapshot: PersonalSnapshot,
        activities: list[Activity],
        output_path: str | Path = "reports/eidon_dashboard.xlsx",
    ) -> Path:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        dashboard = workbook.active
        dashboard.title = "Dashboard"
        dashboard["A1"] = "Eidon OS - Personal Intelligence"
        dashboard["A1"].font = Font(size=16, bold=True)

        metrics = [
            ("Projects", snapshot.projects),
            ("Files", snapshot.files),
            ("People", snapshot.people),
            ("Companies", snapshot.companies),
            ("Knowledge", snapshot.knowledge_items),
            ("Relationships", snapshot.relationships),
            ("Activities", snapshot.activities),
            ("Minutes", snapshot.total_minutes),
        ]
        for row, (label, value) in enumerate(metrics, start=3):
            dashboard.cell(row=row, column=1, value=label)
            dashboard.cell(row=row, column=2, value=value)

        categories = workbook.create_sheet("Categories")
        categories.append(["Category", "Count"])
        for category, count in sorted(snapshot.activity_categories.items()):
            categories.append([category, count])

        if categories.max_row > 1:
            chart = BarChart()
            chart.title = "Activities by category"
            chart.y_axis.title = "Count"
            chart.x_axis.title = "Category"
            data = Reference(categories, min_col=2, min_row=1, max_row=categories.max_row)
            labels = Reference(categories, min_col=1, min_row=2, max_row=categories.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(labels)
            dashboard.add_chart(chart, "D3")

        timeline = workbook.create_sheet("Timeline")
        timeline.append(["Date", "Title", "Category", "Duration (min)", "Description"])
        for activity in activities:
            timeline.append([
                activity.occurred_at.isoformat(),
                activity.title,
                activity.category,
                activity.duration_minutes,
                activity.description,
            ])

        workbook.save(path)
        return path
