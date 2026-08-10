from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from eidon_os.application.personal_intelligence import PersonalSnapshot
from eidon_os.domain.activity import Activity
from eidon_os.domain.goal import Goal
from eidon_os.infrastructure.excel.data_builders import GeneralRow, build_general_rows, filter_by_area


class ExcelWorkbookBuilder:
    AREAS = ("Estudos", "Projetos", "Academia de Código", "Carreira")

    def build(self, snapshot: PersonalSnapshot, activities: list[Activity], output_path: str | Path, *, goals: list[Goal] | None = None) -> Path:
        path = Path(output_path); path.parent.mkdir(parents=True, exist_ok=True); rows = build_general_rows(activities)
        workbook = Workbook(); dashboard = workbook.active; dashboard.title = "Dashboard Geral"
        self._build_dashboard(dashboard, snapshot, rows); self._build_goals(workbook, goals or [])
        for area in self.AREAS: self._build_area_sheet(workbook, area, filter_by_area(rows, area))
        self._build_general_base(workbook, rows); self._build_settings(workbook); self._build_instructions(workbook); workbook.save(path); return path

    def _build_dashboard(self, sheet, snapshot: PersonalSnapshot, rows: list[GeneralRow]) -> None:
        sheet["A1"] = "Eidon OS — Dashboard Geral"; sheet["A1"].font = Font(size=18, bold=True); sheet["A2"] = "Visão consolidada do Personal Intelligence Core."
        metrics = [("Projetos cadastrados", snapshot.projects),("Conhecimentos", snapshot.knowledge_items),("Atividades", snapshot.activities),("Minutos registrados", snapshot.total_minutes),("Horas registradas", round(snapshot.total_minutes/60,1)),("Dias ativos", len({row.date.date() for row in rows}))]
        for i,(label,value) in enumerate(metrics,start=4): sheet.cell(i,1,label); sheet.cell(i,2,value)
        counts=Counter(row.area for row in rows); sheet["D3"]="Área"; sheet["E3"]="Atividades"
        for i,area in enumerate(self.AREAS,start=4): sheet.cell(i,4,area); sheet.cell(i,5,counts.get(area,0))
        chart=BarChart(); chart.title="Atividades por área"; chart.y_axis.title="Quantidade"; chart.add_data(Reference(sheet,min_col=5,min_row=3,max_row=3+len(self.AREAS)),titles_from_data=True); chart.set_categories(Reference(sheet,min_col=4,min_row=4,max_row=3+len(self.AREAS))); sheet.add_chart(chart,"G3"); self._format_header(sheet,3,4,5); sheet.column_dimensions["A"].width=24; sheet.column_dimensions["B"].width=18

    def _build_goals(self, workbook: Workbook, goals: list[Goal]) -> None:
        sheet=workbook.create_sheet("Metas"); sheet.append(["Área","Indicador","Meta","Unidade","Período","Realizado","Progresso"])
        if not goals:
            defaults=[("Academia de Código","Desafios concluídos",40,"desafios","monthly"),("Academia de Código","Horas de estudo",25,"horas","monthly"),("Academia de Código","Plataformas utilizadas",4,"plataformas","monthly"),("Academia de Código","Tecnologias praticadas",3,"tecnologias","monthly")]
            for area,indicator,target,unit,period in defaults: sheet.append([area,indicator,target,unit,period,0,None])
        else:
            for goal in goals: sheet.append([goal.area,goal.indicator,goal.target,goal.unit,goal.period,0,None])
        for row in range(2,sheet.max_row+1): sheet.cell(row,7,f'=IFERROR(F{row}/C{row},0)'); sheet.cell(row,7).number_format="0%"
        self._format_header(sheet,1,1,7); self._add_table(sheet,"MetasTable"); self._fit_columns(sheet)

    def _build_area_sheet(self,workbook:Workbook,area:str,rows:list[GeneralRow])->None:
        sheet=workbook.create_sheet(area[:31]); sheet.append(["Data","Atividade","Categoria","Minutos","Origem","Status","Descrição"])
        for row in rows: sheet.append([row.date.replace(tzinfo=None),row.activity,row.category,row.minutes,row.origin,row.status,row.description])
        self._format_header(sheet,1,1,7); self._add_table(sheet,self._table_name(area)); self._fit_columns(sheet)

    def _build_general_base(self,workbook:Workbook,rows:list[GeneralRow])->None:
        sheet=workbook.create_sheet("Base Geral"); sheet.append(["Data","Área","Atividade","Categoria","Minutos","Origem","Status","Descrição"])
        for row in rows: sheet.append([row.date.replace(tzinfo=None),row.area,row.activity,row.category,row.minutes,row.origin,row.status,row.description])
        self._format_header(sheet,1,1,8); self._add_table(sheet,"BaseGeralTable"); self._fit_columns(sheet)

    def _build_settings(self,workbook:Workbook)->None:
        sheet=workbook.create_sheet("Listas e Configurações"); sheet.append(["Áreas","Status"]); statuses=("Planejado","Em andamento","Concluído","Pausado")
        for i in range(max(len(self.AREAS),len(statuses))): sheet.append([self.AREAS[i] if i<len(self.AREAS) else None,statuses[i] if i<len(statuses) else None])
        self._format_header(sheet,1,1,2); self._fit_columns(sheet)

    def _build_instructions(self,workbook:Workbook)->None:
        sheet=workbook.create_sheet("Instruções"); sheet["A1"]="Como usar o Dashboard"; sheet["A1"].font=Font(size=16,bold=True)
        items=["O Dashboard Geral é somente visualização.","Os dados principais vivem no SQLite do Eidon OS.","Cada área possui sua própria visualização.","A Base Geral consolida os campos comuns.","As metas cadastradas no Eidon alimentam a aba Metas.","Excel é formato de exportação; a interface principal será o Eidon Web na v0.5.0.","Execute 'eidon dashboard-export' para regenerar o arquivo."]
        for i,text in enumerate(items,start=3): sheet.cell(i,1,f"{i-2}. {text}")
        sheet.column_dimensions["A"].width=110

    @staticmethod
    def _format_header(sheet,row:int,start_column:int,end_column:int)->None:
        fill=PatternFill("solid",fgColor="1F4E78")
        for column in range(start_column,end_column+1):
            cell=sheet.cell(row,column); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=fill; cell.alignment=Alignment(horizontal="center")
    @staticmethod
    def _add_table(sheet,name:str)->None:
        if sheet.max_row<2: sheet.append([None]*sheet.max_column)
        table=Table(displayName=name,ref=f"A1:{sheet.cell(sheet.max_row,sheet.max_column).coordinate}"); table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False); sheet.add_table(table)
    @staticmethod
    def _fit_columns(sheet)->None:
        for cells in sheet.columns: sheet.column_dimensions[cells[0].column_letter].width=min(max(len(str(cell.value or "")) for cell in cells)+2,45)
    @staticmethod
    def _table_name(area:str)->str:
        normalized="".join(ch for ch in area if ch.isalnum()); return f"{normalized}Table"[:255]
